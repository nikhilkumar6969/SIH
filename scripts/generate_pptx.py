import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set slide dimensions to 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_BG_DARK = RGBColor(15, 23, 42)      # #0F172A Slate 900
    COLOR_BG_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Slate 50
    COLOR_CARD_BG = RGBColor(255, 255, 255)   # #FFFFFF
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 Slate 200

    COLOR_PRIMARY = RGBColor(37, 99, 235)     # #2563EB Blue 600
    COLOR_PRIMARY_DARK = RGBColor(29, 78, 216) # #1D4ED8 Blue 700
    COLOR_TEAL = RGBColor(13, 148, 136)       # #0D9488 Teal 600
    COLOR_PURPLE = RGBColor(124, 58, 237)     # #7C3AED Purple 600
    COLOR_AMBER = RGBColor(217, 119, 6)       # #D97706 Amber 600
    COLOR_GREEN = RGBColor(22, 163, 74)       # #16A34A Green 600
    
    COLOR_TEXT_DARK = RGBColor(15, 23, 42)    # Slate 900
    COLOR_TEXT_MUTED = RGBColor(71, 85, 105)  # Slate 600
    COLOR_TEXT_LIGHT = RGBColor(241, 245, 249)# Slate 100
    COLOR_TEXT_SUBTLE = RGBColor(148, 163, 184)# Slate 400

    FONT_MAIN = "Arial"
    FONT_HEADING = "Arial"

    def add_header(slide, title, category="SMART INDIA HACKATHON 2026"):
        # Header category tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.font.name = FONT_MAIN

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK
        p.font.name = FONT_HEADING

    def add_footer(slide, current_page, total_pages=8):
        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CARD_BORDER
        line.line.color.rgb = COLOR_CARD_BORDER

        # Footer Text
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(10.0), Inches(0.3))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "SIH 2026 | Problem Statement ID: 26034 | Team: NEW COMERS (SIH26_CGITJDP_001)"
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.font.name = FONT_MAIN

        # Page Number
        page_box = slide.shapes.add_textbox(Inches(11.5), Inches(7.0), Inches(1.0), Inches(0.3))
        tf = page_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{current_page} / {total_pages}"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.font.name = FONT_MAIN

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: PROJECT INTRODUCTION
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_BG_DARK)

    # Decorative header shape
    dec_shape = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(1.2), Inches(0.06))
    dec_shape.fill.solid()
    dec_shape.fill.fore_color.rgb = COLOR_PRIMARY
    dec_shape.line.fill.background()

    # Title & Subtitle Container
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.7), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "LABELGUARD"
    p0.font.size = Pt(40)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEXT_LIGHT
    p0.font.name = FONT_HEADING

    p1 = tf.add_paragraph()
    p1.text = "AI-Powered Statutory Compliance Platform for Legal Metrology Packaged Commodities"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.font.name = FONT_MAIN
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Automated evidentiary OCR detection, 36 statutory rule engine, visual contrast & font height analysis, and official enforcement notice generation."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_TEXT_SUBTLE
    p2.font.name = FONT_MAIN
    p2.space_before = Pt(10)

    # Card 1: Problem Statement Details
    card1 = add_card(slide1, 0.8, 3.7, 5.7, 2.8, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(5.3), Inches(2.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "PROBLEM STATEMENT DETAILS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    items1 = [
        ("PS ID:", " 26034"),
        ("PS Title:", " Software System to check compliance of Packaged Commodities under Legal Metrology Rules, 2011 by scanning products, images and labels"),
        ("Theme:", " Miscellaneous"),
        ("Category:", " Software")
    ]
    for label, val in items1:
        p = tf1.add_paragraph()
        p.space_before = Pt(6)
        run1 = p.add_run()
        run1.text = label
        run1.font.bold = True
        run1.font.size = Pt(10)
        run1.font.color.rgb = COLOR_TEXT_LIGHT
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(10)
        run2.font.color.rgb = COLOR_TEXT_SUBTLE

    # Card 2: Team Details
    card2 = add_card(slide1, 6.8, 3.7, 5.7, 2.8, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    tb2 = slide1.shapes.add_textbox(Inches(7.0), Inches(3.9), Inches(5.3), Inches(2.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "TEAM & REGISTRATION INFO"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE

    items2 = [
        ("Hackathon:", " Smart India Hackathon (SIH) 2026"),
        ("Team Name:", " NEW COMERS"),
        ("Team ID:", " SIH26_CGITJDP_001"),
        ("Target Domain:", " Legal Metrology Act 2009 & Packaged Commodities Rules 2011"),
        ("Core Stack:", " React, TypeScript, FastAPI, RapidOCR (ONNX), Supabase")
    ]
    for label, val in items2:
        p = tf2.add_paragraph()
        p.space_before = Pt(6)
        run1 = p.add_run()
        run1.text = label
        run1.font.bold = True
        run1.font.size = Pt(10)
        run1.font.color.rgb = COLOR_TEXT_LIGHT
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(10)
        run2.font.color.rgb = COLOR_TEXT_SUBTLE

    # Footer for Slide 1
    f_box = slide1.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.3))
    tf = f_box.text_frame
    p = tf.paragraphs[0]
    p.text = "CONFIDENTIAL & PROPRIETARY — PREPARED FOR SMART INDIA HACKATHON 2026 EVALUATION"
    p.font.size = Pt(9)
    p.font.color.rgb = COLOR_TEXT_SUBTLE

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT & PROPOSED SOLUTION
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_BG_LIGHT)
    add_header(slide2, "PROBLEM STATEMENT & PROPOSED SOLUTION")
    add_footer(slide2, 2)

    # Left Column: Real-World Problem
    add_card(slide2, 0.8, 1.4, 5.6, 5.3)
    tb_p = slide2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(4.9))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    p.text = "REAL-WORLD PROBLEM & CHALLENGES"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(225, 29, 72) # Red/Rose accent

    prob_points = [
        ("Manual & Slow Inspections: ", "Enforcement officers must physically inspect thousands of packaged commodities across physical stores and digital marketplaces."),
        ("Deceptive & Missing Declarations: ", "Key statutory details (MRP inclusive taxes, Unit Sale Price, Country of Origin, Expiry, Manufacturer Address) are frequently missing or altered."),
        ("Legibility & Font Height Violations: ", "Rule 13 mandates minimum physical font height (1.6mm) and contrast standards that are impossible to verify accurately by human eye alone."),
        ("E-Commerce Compliance Blindspots: ", "Rule 6(10) mandates full statutory declarations on e-commerce PDPs (Amazon, Flipkart, Blinkit, etc.) where violations go unchecked at scale."),
        ("No Traceable Audit Trail: ", "Paper-based notices lead to legal disputes, slow court enforcement, and lack of tamper-proof evidentiary logs.")
    ]
    for title, desc in prob_points:
        p = tf_p.add_paragraph()
        p.space_before = Pt(8)
        run1 = p.add_run()
        run1.text = title
        run1.font.bold = True
        run1.font.size = Pt(10)
        run1.font.color.rgb = COLOR_TEXT_DARK
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Column: Proposed Solution
    add_card(slide2, 6.8, 1.4, 5.7, 5.3)
    tb_s = slide2.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(4.9))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "THE LABELGUARD SOLUTION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    sol_points = [
        ("Deep Learning OCR Extraction: ", "RapidOCR (PP-OCRv4 DBNet + SVTR) microservice extracting text, bounding boxes, and confidence scores in <2.3s."),
        ("36-Rule Statutory Compliance Engine: ", "Automated evaluation of 28 physical package rules + 8 e-commerce PDP rules under Legal Metrology Rules, 2011."),
        ("Computer Vision Font & Contrast Analysis: ", "Spatial pixel sampling estimating font height in mm against Rule 13 baselines + WCAG luminance contrast checking."),
        ("Human-in-the-Loop Determination: ", "Strict legal safety architecture—officers review, modify, or dismiss AI findings before issuing formal notices."),
        ("Multi-Format Legal Dossier Generation: ", "Instant creation of official PDF Inspection Certificates and editable DOCX statutory legal notices under LM Act Sec 36.")
    ]
    for title, desc in sol_points:
        p = tf_s.add_paragraph()
        p.space_before = Pt(8)
        run1 = p.add_run()
        run1.text = title
        run1.font.bold = True
        run1.font.size = Pt(10)
        run1.font.color.rgb = COLOR_TEXT_DARK
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 3: UNIQUE VALUE PROPOSITION & INNOVATION
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_BG_LIGHT)
    add_header(slide3, "UNIQUE VALUE PROPOSITION & TECHNICAL INNOVATION")
    add_footer(slide3, 3)

    uvp_cards = [
        ("Legal Safety & Ethical AI Architecture", 
         "Non-autonomous design: LabelGuard functions strictly as an evidentiary aid. Every AI detection includes spatial bounding box evidence, rule references, confidence ratings, and mandatory officer endorsement controls.", 
         COLOR_PRIMARY, 0.8, 1.4, 5.6, 2.5),
        
        ("Dual-Mode Edge & Cloud Deployment", 
         "Runs seamlessly online via Supabase Cloud (PostgreSQL, Auth, Storage) OR fully offline in Local Mock Mode for remote field enforcement officers in low-connectivity retail inspection areas.", 
         COLOR_TEAL, 6.8, 1.4, 5.7, 2.5),
        
        ("Multi-Channel Physical & Digital Audit", 
         "Unified compliance platform auditing both physical retail packaging (camera photo scan) AND digital e-commerce product listings across Amazon, Flipkart, Blinkit, Zepto, JioMart, and BigBasket.", 
         COLOR_PURPLE, 0.8, 4.2, 5.6, 2.5),
        
        ("Automated Legal Dossier & DOCX Drafting", 
         "Generates official tamper-evident PDF inspection dossiers AND fully editable Microsoft Word (.docx) formal statutory non-compliance notices under Section 36 of the Legal Metrology Act, 2009.", 
         COLOR_AMBER, 6.8, 4.2, 5.7, 2.5)
    ]

    for title, desc, color, x, y, w, h in uvp_cards:
        add_card(slide3, x, y, w, h)
        # Accent bar on left of card
        bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.12), Inches(h))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = slide3.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.2), Inches(w - 0.5), Inches(h - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.space_before = Pt(6)
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, COLOR_BG_LIGHT)
    add_header(slide4, "SYSTEM ARCHITECTURE & DATA FLOW")
    add_footer(slide4, 4)

    # Architecture Columns (4 Layers)
    arch_layers = [
        ("1. INPUT LAYER", ["Camera Photo Capture", "Label Image Upload", "E-Commerce PDP Scraper", "Batch Image Ingestion"], COLOR_PRIMARY, 0.8),
        ("2. OCR & VISION PIPELINE", ["RapidOCR (PP-OCRv4 DBNet)", "ONNX Runtime Inference", "Smart 4-Way Rotation Detection", "Bounding Box & Spatial Extent"], COLOR_TEAL, 3.8),
        ("3. STATUTORY ENGINE", ["36-Rule Compliance Engine", "Font Height Estimator (mm)", "WCAG Contrast Ratio Analyzer", "Rule 26 Medical Exemption Router"], COLOR_PURPLE, 6.8),
        ("4. OUTPUT & DOSSIER", ["Officer Determination Workspace", "Finding Overrides & Audit Log", "Official PDF Dossier Export", "Editable DOCX Notice Generator"], COLOR_AMBER, 9.8)
    ]

    for title, items, color, x in arch_layers:
        add_card(slide4, x, 1.4, 2.7, 4.3)
        # Header banner inside card
        header_shape = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.4), Inches(2.7), Inches(0.5))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = color
        header_shape.line.fill.background()

        tb_h = slide4.shapes.add_textbox(Inches(x + 0.1), Inches(1.45), Inches(2.5), Inches(0.4))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True
        p = tf_h.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_BG

        tb_c = slide4.shapes.add_textbox(Inches(x + 0.2), Inches(2.0), Inches(2.3), Inches(3.6))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        for idx, item in enumerate(items):
            p = tf_c.paragraphs[0] if idx == 0 else tf_c.add_paragraph()
            if idx > 0:
                p.space_before = Pt(8)
            p.text = f"•  {item}"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Data Flow Bar
    add_card(slide4, 0.8, 5.9, 11.7, 0.8, bg_color=RGBColor(241, 245, 249), border_color=COLOR_PRIMARY)
    tb_df = slide4.shapes.add_textbox(Inches(1.0), Inches(5.95), Inches(11.3), Inches(0.7))
    tf_df = tb_df.text_frame
    tf_df.word_wrap = True

    p = tf_df.paragraphs[0]
    p.text = "DATA STORAGE & AUTHENTICATION LAYER:"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p2 = tf_df.add_paragraph()
    p2.text = "Supabase PostgreSQL Database (Row-Level Security)  ↔  Supabase Auth & Storage  ↔  Local Mock Mode Fallback for Offline Edge Devices"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 5: TECHNICAL APPROACH & TECHNOLOGY STACK
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, COLOR_BG_LIGHT)
    add_header(slide5, "TECHNICAL WORKFLOW & STACK IMPLEMENTATION")
    add_footer(slide5, 5)

    # Left Box: End-to-End Processing Workflow
    add_card(slide5, 0.8, 1.4, 6.2, 5.3)
    tb_wf = slide5.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.8), Inches(4.9))
    tf_wf = tb_wf.text_frame
    tf_wf.word_wrap = True

    p = tf_wf.paragraphs[0]
    p.text = "6-STEP END-TO-END PROCESSING PIPELINE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    steps = [
        ("Step 1: Image Capture & Preprocessing", "Image upload or camera scan; quality check for contrast, resolution, and blur."),
        ("Step 2: RapidOCR ONNX Recognition", "FastAPI microservice runs PP-OCRv4 DBNet + SVTR with smart 4-orientation rotation scan."),
        ("Step 3: Regex Normalization", "Cleans extracted text for currency (₹), dates (MM/YYYY), SI units (g, kg, ml), and FSSAI license numbers."),
        ("Step 4: 36-Rule Statutory Evaluation", "Evaluates mandatory declarations, Unit Sale Price (USP), font heights (mm), and WCAG contrast ratio."),
        ("Step 5: Human Officer Verification Workspace", "Flags visually ambiguous or subjective rules for officer determination; supports individual rule overrides."),
        ("Step 6: Statutory Evidence & Dossier Export", "Generates immutable inspection log + official PDF certificate + editable DOCX formal legal notice.")
    ]
    for title, desc in steps:
        p = tf_wf.add_paragraph()
        p.space_before = Pt(6)
        run1 = p.add_run()
        run1.text = title + " — "
        run1.font.bold = True
        run1.font.size = Pt(9.5)
        run1.font.color.rgb = COLOR_TEXT_DARK
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(9)
        run2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Box: Complete Technology Stack
    add_card(slide5, 7.3, 1.4, 5.2, 5.3)
    tb_st = slide5.shapes.add_textbox(Inches(7.5), Inches(1.6), Inches(4.8), Inches(4.9))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True

    p = tf_st.paragraphs[0]
    p.text = "VERIFIED TECHNOLOGY STACK"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    stack_items = [
        ("Frontend Framework:", " React 18, TypeScript, Vite, Tailwind CSS"),
        ("Icons & UI Components:", " Lucide React, Recharts analytics gauges"),
        ("OCR Backend Microservice:", " FastAPI (Python 3.11/3.12), RapidOCR"),
        ("Inference Engine:", " ONNX Runtime, OpenCV, PIL, NumPy"),
        ("Database & Auth:", " Supabase (PostgreSQL, RLS, Storage)"),
        ("Offline Edge Fallback:", " Local Mock Mode state container"),
        ("Document Generation:", " jsPDF & AutoTable (PDF), docx.js (.docx)")
    ]
    for label, val in stack_items:
        p = tf_st.add_paragraph()
        p.space_before = Pt(7)
        run1 = p.add_run()
        run1.text = label
        run1.font.bold = True
        run1.font.size = Pt(9.5)
        run1.font.color.rgb = COLOR_TEXT_DARK
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 6: FEASIBILITY, VIABILITY & RISK MITIGATION
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, COLOR_BG_LIGHT)
    add_header(slide6, "FEASIBILITY, VIABILITY & RISK MITIGATION")
    add_footer(slide6, 6)

    # 3 Cards Layout: Feasibility, Viability, Risks
    card_data = [
        ("TECHNICAL & OPERATIONAL FEASIBILITY", COLOR_PRIMARY, [
            ("Proven OCR Foundation: ", "Built on RapidOCR/PaddleOCR, an open-source deep learning model validated across complex multilingual product labels."),
            ("Dual Mode Architecture: ", "Runs seamlessly on Supabase Cloud or offline edge devices for field inspections in low-network regions."),
            ("Configurable 36-Rule Engine: ", "Easily extends to new commodity categories and statutory updates without code changes.")
        ], 0.8, 1.4, 3.6, 5.3),

        ("COMMERCIAL & SYSTEM VIABILITY", COLOR_TEAL, [
            ("Government Cost & Effort Savings: ", "Cuts manual label verification time from minutes to seconds across hundreds of product categories."),
            ("E-Commerce Scale Ready: ", "Pre-built auditing for Amazon, Flipkart, Blinkit, Zepto, JioMart, and BigBasket PDP listings."),
            ("Statutory Notice Efficiency: ", "Auto-generated editable DOCX notices reduce legal team drafting time by over 80%.")
        ], 4.8, 1.4, 3.7, 5.3),

        ("RISKS & PRACTICAL MITIGATION", COLOR_AMBER, [
            ("Risk: AI Legal Misattribution", "Mitigation: Strict Human-in-the-Loop policy. AI findings act purely as evidence; final decision rests with certified officers."),
            ("Risk: Poor Photo Quality / Rotation", "Mitigation: Smart 4-way multi-orientation OCR scan + image resolution & blur preprocessing checks."),
            ("Risk: E-Commerce Layout Changes", "Mitigation: Normalization regex & DOM abstraction layers separating listing data from visual OCR.")
        ], 8.8, 1.4, 3.7, 5.3)
    ]

    for title, color, items, x, y, w, h in card_data:
        add_card(slide6, x, y, w, h)
        bar = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb_h = slide6.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.05), Inches(w - 0.2), Inches(0.35))
        tf_h = tb_h.text_frame
        p = tf_h.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_BG

        tb_c = slide6.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.5), Inches(w - 0.4), Inches(h - 0.6))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        for idx, (head, desc) in enumerate(items):
            p = tf_c.paragraphs[0] if idx == 0 else tf_c.add_paragraph()
            if idx > 0:
                p.space_before = Pt(8)
            run1 = p.add_run()
            run1.text = head + " "
            run1.font.bold = True
            run1.font.size = Pt(9)
            run1.font.color.rgb = COLOR_TEXT_DARK
            run2 = p.add_run()
            run2.text = desc
            run2.font.size = Pt(8.5)
            run2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 7: IMPACT, BENEFITS & FUTURE SCOPE
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, COLOR_BG_LIGHT)
    add_header(slide7, "IMPACT, USER BENEFITS & ROADMAP")
    add_footer(slide7, 7)

    # Top Row: Impact Cards
    impact_items = [
        ("ENFORCEMENT EFFICIENCY", "Transforms slow paper inspections into automated evidentiary scans, scaling enforcement coverage 10x.", COLOR_PRIMARY, 0.8, 1.4, 3.6, 2.4),
        ("CONSUMER PROTECTION", "Prevents illegal price markups, hidden taxes, misleading unit sale prices, and non-compliant packaging.", COLOR_TEAL, 4.8, 1.4, 3.7, 2.4),
        ("E-COMMERCE COMPLIANCE", "Enforces Rule 6(10) statutory declarations across digital marketplaces, protecting online shoppers.", COLOR_PURPLE, 8.8, 1.4, 3.7, 2.4)
    ]

    for title, desc, color, x, y, w, h in impact_items:
        add_card(slide7, x, y, w, h)
        bar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.12), Inches(h))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = slide7.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.2), Inches(w - 0.4), Inches(h - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.space_before = Pt(6)
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Container: Future Development Roadmap
    add_card(slide7, 0.8, 4.1, 11.7, 2.6)
    tb_rm = slide7.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(11.3), Inches(2.3))
    tf_rm = tb_rm.text_frame
    tf_rm.word_wrap = True

    p = tf_rm.paragraphs[0]
    p.text = "FUTURE DEVELOPMENT ROADMAP & PLANNED EXTENSIONS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    roadmap_points = [
        ("Phase 1 — Physical Calibration Target Integration: ", "Adding millimeter scale ruler cards in field photos to improve font height estimation precision."),
        ("Phase 2 — VLM (Visual Language Model) Symbol Detection: ", "Integrating multimodal vision models to detect statutory symbols (Vegetarian/Non-Veg dot, Recycling, ISO marks)."),
        ("Phase 3 — Multi-Tenant State Legal Metrology RLS: ", "Implementing state-wise Row Level Security for multi-jurisdiction enforcement agencies across India."),
        ("Phase 4 — Automated Marketplace Notice Dispatch: ", "Direct API integration with seller portals for issuing electronic Section 36 statutory non-compliance notices.")
    ]
    for title, desc in roadmap_points:
        p = tf_rm.add_paragraph()
        p.space_before = Pt(5)
        run1 = p.add_run()
        run1.text = title
        run1.font.bold = True
        run1.font.size = Pt(9.5)
        run1.font.color.rgb = COLOR_TEXT_DARK
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(9)
        run2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 8: CONCLUSION, DEMO & REFERENCES
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, COLOR_BG_DARK)

    # Title & Conclusion Header
    t_box8 = slide8.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.5))
    tf8 = t_box8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "LABELGUARD: READY FOR DEPLOYMENT"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_LIGHT

    p2 = tf8.add_paragraph()
    p2.text = "A robust, legally sound, AI-powered compliance platform empowering Legal Metrology officers to uphold consumer rights across physical and digital markets."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_PRIMARY
    p2.space_before = Pt(6)

    # Card 1: Implementation Status
    card8_1 = add_card(slide8, 0.8, 2.5, 5.6, 4.0, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    tb8_1 = slide8.shapes.add_textbox(Inches(1.0), Inches(2.7), Inches(5.2), Inches(3.6))
    tf8_1 = tb8_1.text_frame
    tf8_1.word_wrap = True

    p = tf8_1.paragraphs[0]
    p.text = "CURRENT IMPLEMENTATION STATUS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    status_items = [
        ("React + Vite Frontend:", " 100% Operational"),
        ("FastAPI + RapidOCR Microservice:", " Fully Functional (<2.3s latency)"),
        ("36-Rule Compliance Engine:", " Active & Tested"),
        ("PDF & DOCX Export Engines:", " Functional"),
        ("Supabase + Local Mock Mode:", " Integrated")
    ]
    for label, val in status_items:
        p = tf8_1.add_paragraph()
        p.space_before = Pt(8)
        run1 = p.add_run()
        run1.text = "✓  " + label
        run1.font.bold = True
        run1.font.size = Pt(10)
        run1.font.color.rgb = COLOR_TEXT_LIGHT
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = COLOR_TEXT_SUBTLE

    # Card 2: References & Statutory Basis
    card8_2 = add_card(slide8, 6.8, 2.5, 5.7, 4.0, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    tb8_2 = slide8.shapes.add_textbox(Inches(7.0), Inches(2.7), Inches(5.3), Inches(3.6))
    tf8_2 = tb8_2.text_frame
    tf8_2.word_wrap = True

    p = tf8_2.paragraphs[0]
    p.text = "STATUTORY REFERENCES & DATA SOURCES"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    ref_items = [
        ("Legal Metrology Act, 2009:", " Sections 11, 12, 18, and 36 statutory provisions"),
        ("LM (Packaged Commodities) Rules, 2011:", " Rules 6(1), 6(10), 6(11), 7, 9, 12, 13 (Font Height Table 1)"),
        ("FSSAI Packaging & Labelling Norms:", " Licensing & nutritional panel standards"),
        ("Deep Learning Models:", " PaddleOCR / RapidOCR (PP-OCRv4 DBNet + SVTR ONNX Runtime)"),
        ("Web Content Accessibility Guidelines:", " WCAG 2.1 Luminance Contrast Ratio Standards")
    ]
    for label, val in ref_items:
        p = tf8_2.add_paragraph()
        p.space_before = Pt(8)
        run1 = p.add_run()
        run1.text = "•  " + label
        run1.font.bold = True
        run1.font.size = Pt(9.5)
        run1.font.color.rgb = COLOR_TEXT_LIGHT
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(9)
        run2.font.color.rgb = COLOR_TEXT_SUBTLE

    # Footer Slide 8
    f_box8 = slide8.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.3))
    tf8_f = f_box8.text_frame
    p = tf8_f.paragraphs[0]
    p.text = "THANK YOU — TEAM NEW COMERS (SIH26_CGITJDP_001) | SMART INDIA HACKATHON 2026"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    # Save presentation
    output_path = os.path.join(os.getcwd(), "LabelGuard_SIH2026_Pitch_Deck.pptx")
    prs.save(output_path)
    print(f"[+] PPTX successfully generated at: {output_path}")

if __name__ == "__main__":
    create_deck()
